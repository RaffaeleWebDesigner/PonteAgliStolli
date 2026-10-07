#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Crea il file Excel "modello" da importare in Google Fogli, gia' riempito con i dati attuali del sito.

    python tools/crea_modello_foglio.py [percorso/del/file.xlsx]

Il foglio contiene:
  - Partite, Reti, Pagelle  → le schede che aggiornano gli altri (risultati, marcatori, pagelle)
  - Rosa, Squadre           → elenchi di riferimento per i menu a tendina (li gestisce il proprietario del foglio)
Classifica, staff, sponsor e campi NON sono nel foglio: si modificano in assets/js/data.js.

Serve solo la prima volta (o per rifare il modello / aggiornare gli elenchi). Richiede: pip install openpyxl
"""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RADICE = Path(__file__).resolve().parent.parent
DESTINAZIONE = Path(sys.argv[1]) if len(sys.argv) > 1 else RADICE / "Modello foglio ASD Ponte agli Stolli.xlsx"
SQUADRA = "Ponte agli Stolli"
RIGHE_UTILI = 400      # righe pre-formattate (testo per date e ore)

BORDEAUX = "6D1B34"
GRIGIO = "808080"

# ---- legge i dati attuali del sito da data.js
codice = ("const vm=require('vm'),fs=require('fs');const c=fs.readFileSync(process.argv[1],'utf8');"
          "const x={};vm.createContext(x);vm.runInContext(c+';this.__d=SITE_DATA;',x);console.log(JSON.stringify(x.__d))")
dati = json.loads(subprocess.run(["node", "-e", codice, str(RADICE / "assets" / "js" / "data.js")],
                                 capture_output=True, text=True, check=True, encoding="utf-8").stdout)


def data_it(iso):
    if not iso:
        return ""
    a, m, g = iso.split("-")
    return f"{g}/{m}/{a}"


def etichetta(nome, numero):
    return f"{nome} #{numero}" if numero is not None else nome


rosa = dati["giocatori"]
per_nome = {}
for p in rosa:
    per_nome.setdefault(p["nome"], []).append(p)

wb = Workbook()
wb.remove(wb.active)


def scheda(nome, intestazioni, larghezze, righe, testo_colonne=(), colore=BORDEAUX):
    ws = wb.create_sheet(nome)
    ws.append(intestazioni)
    for r in righe:
        ws.append(r)
    for i, _ in enumerate(intestazioni, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=colore)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = larghezze[i - 1]
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    for col in testo_colonne:           # le date e le ore restano testo: niente conversioni automatiche
        for r in range(2, RIGHE_UTILI + 2):
            ws.cell(row=r, column=col).number_format = "@"
    return ws


def tendina(ws, colonna, formula, righe_fino=RIGHE_UTILI + 1, rigida=False, messaggio=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                        errorStyle="stop" if rigida else "warning",
                        errorTitle="Valore non in elenco", error=messaggio or "Questo valore non e' nell'elenco: controlla di averlo scritto bene.")
    ws.add_data_validation(dv)
    dv.add(f"{colonna}2:{colonna}{righe_fino}")


# =============================================================== Istruzioni
ws = wb.create_sheet("Istruzioni")
testi = [
    ("COME SI AGGIORNANO I RISULTATI SUL SITO DELLA ASD PONTE AGLI STOLLI", "titolo"),
    ("Il sito legge questo foglio ogni 15 minuti circa: quello che scrivi nelle schede Partite, Reti e Pagelle compare sul sito.", None),
    ("La classifica, la rosa, lo staff e gli sponsor NON si cambiano da qui: li aggiorna Raffaele.", None),
    ("", None),
    ("REGOLE GENERALI", "sezione"),
    ("• Non cambiare i nomi delle schede (le linguette in basso) e non cambiare la prima riga di ogni scheda.", None),
    ("• Date: scrivile come GG/MM/AAAA, ad esempio 09/10/2026. Ora: 21:00.", None),
    ("• Imposta il foglio in italiano: File > Impostazioni > Impostazioni internazionali: Italia.", None),
    ("• Dove c'e' una freccina nella cella, scegli il valore dall'elenco invece di scriverlo (evita errori di battitura).", None),
    ("• Le schede Rosa e Squadre servono solo per gli elenchi a tendina: non modificarle.", None),
    ("", None),
    ("DOPO UNA PARTITA", "sezione"),
    ("1. Scheda Partite: nella riga della partita scrivi «Gol casa» e «Gol ospite». La partita passa da sola da Calendario a Risultati.", None),
    ("2. Scheda Reti: una riga per ogni marcatore → Data della partita, Avversario, Giocatore (dall'elenco), Gol (solo se ne ha fatti piu' di uno).", None),
    ("3. (Facoltativo) Scheda Pagelle: una riga per giocatore → Data, Avversario, Giocatore, Voto, Commento.", None),
    ("I marcatori della stagione (solo Campionato e Coppa, non le amichevoli) e i gol dei giocatori si calcolano da soli dalla scheda Reti.", None),
    ("", None),
    ("QUANDO SI CONOSCONO DATA, ORA E CAMPO DI UNA PARTITA", "sezione"),
    ("Scheda Partite: compila Data, Ora e Luogo nella riga della partita (il Luogo va scritto esattamente come nelle altre partite in casa).", None),
    ("", None),
    ("SE QUALCOSA NON FUNZIONA", "sezione"),
    ("Se un dato e' sbagliato (una data che non esiste, un giocatore non in elenco, marcatori piu' numerosi dei gol del risultato...) il sito NON si aggiorna e resta com'era.", None),
    ("Se dopo mezz'ora la modifica non e' sul sito, ricontrolla quello che hai scritto. Se non trovi l'errore, scrivi a Raffaele: gli arriva l'indicazione esatta della riga da correggere.", None),
    ("Se un giocatore e' nuovo o ha cambiato numero, avvisa Raffaele: aggiorna lui rosa ed elenchi.", None),
]
for i, (t, stile) in enumerate(testi, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if stile == "titolo":
        c.font = Font(bold=True, size=15, color=BORDEAUX)
    elif stile == "sezione":
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BORDEAUX)
ws.column_dimensions["A"].width = 130

# =============================================================== Partite (calendario + risultati in un'unica scheda)
righe = []
for r in dati["risultati"]:
    righe.append([r["competizione"], r["giornata"] if r["giornata"] is not None else "", data_it(r["data"]), "", "",
                  r["casa"], r["ospite"], "", r["golCasa"], r["golOspite"]])
for chiave, nome in (("campionato", "Campionato"), ("coppa", "Coppa"), ("amichevoli", "Amichevole")):
    for m in dati["calendario"][chiave]:
        righe.append([nome, m["giornata"] if m["giornata"] is not None else "", data_it(m["data"]), m["ora"] or "",
                      m.get("periodo") or "", m["casa"], m["ospite"], m["luogo"] or "", "", ""])
ws = scheda("Partite",
            ["Competizione", "Giornata", "Data", "Ora", "Periodo", "Casa", "Ospite", "Luogo", "Gol casa", "Gol ospite"],
            [15, 14, 13, 8, 28, 26, 26, 46, 10, 11], righe, testo_colonne=(3, 4, 5))
tendina(ws, "A", '"Campionato,Coppa,Amichevole"', rigida=True)
tendina(ws, "F", "=Squadre!$A$2:$A$60")
tendina(ws, "G", "=Squadre!$A$2:$A$60")
dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                    errorTitle="Numero non valido", error="Scrivi un numero intero (0, 1, 2, ...).")
ws.add_data_validation(dv)
dv.add(f"I2:J{RIGHE_UTILI + 1}")

# =============================================================== Reti
righe = []
for r in dati["risultati"]:
    avversario = r["ospite"] if r["casa"] == SQUADRA else r["casa"]
    for x in r["reti"]:
        numero = x.get("numero")
        if numero is None and len(per_nome.get(x["giocatore"], [])) == 1:
            numero = per_nome[x["giocatore"]][0]["numero"]
        righe.append([data_it(r["data"]), avversario, etichetta(x["giocatore"], numero), x["gol"]])
ws = scheda("Reti", ["Data", "Avversario", "Giocatore", "Gol"], [14, 28, 40, 8], righe, testo_colonne=(1,))
tendina(ws, "B", "=Squadre!$A$2:$A$60")
tendina(ws, "C", "=Rosa!$A$2:$A$80", messaggio="Scegli il giocatore dall'elenco (con il numero di maglia). Se e' un giocatore nuovo, avvisa Raffaele.")

# =============================================================== Pagelle
righe = []
for r in dati["risultati"]:
    avversario = r["ospite"] if r["casa"] == SQUADRA else r["casa"]
    for p in r["pagelle"]:
        righe.append([data_it(r["data"]), avversario, p["giocatore"], p["voto"], p["nota"]])
ws = scheda("Pagelle", ["Data", "Avversario", "Giocatore", "Voto", "Commento"], [14, 28, 40, 8, 70], righe, testo_colonne=(1,))
tendina(ws, "B", "=Squadre!$A$2:$A$60")
tendina(ws, "C", "=Rosa!$A$2:$A$80", messaggio="Scegli il giocatore dall'elenco (con il numero di maglia).")

# =============================================================== Rosa e Squadre: SOLO elenchi di riferimento (grigi)
ws = scheda("Rosa", ["Giocatore (NON MODIFICARE)", "Ruolo"], [44, 18],
            [[etichetta(p["nome"], p["numero"]), p["ruolo"]] for p in rosa], colore=GRIGIO)
ws = scheda("Squadre", ["Squadra (NON MODIFICARE)"], [40], [[c["squadra"]] for c in dati["classifica"]], colore=GRIGIO)

wb.save(DESTINAZIONE)
print("Creato:", DESTINAZIONE)
